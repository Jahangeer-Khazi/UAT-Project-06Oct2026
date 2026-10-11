"""Build the Case duplicate-field Excel report from analysis.json."""
import json, re, sys
from collections import defaultdict, Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

d = json.load(open(sys.argv[1]))
OUT = sys.argv[2]
F, R = d['fields'], d['refs']
FUNC, UI, CONF = d['cats']['FUNCTIONAL'], d['cats']['UI'], d['cats']['CONFIG']

# compact reference buckets shown as columns
BUCKETS = [
    ('Apex (Classes/Triggers)', 'F', ['Apex Class', 'Apex Trigger']),
    ('Flows', 'F', ['Flow']),
    ('LWC / Aura / Visualforce', 'F', ['LWC', 'Aura', 'Visualforce']),
    ('Email Templates', 'F', ['Email Template']),
    ('Validation Rules', 'F', ['Validation Rule']),
    ('Formula / Other Fields', 'F', ['Case Formula/Field', 'Cross-Object Formula/Field']),
    ('Workflow / Approval / Assignment Rules', 'F', ['Workflow', 'Approval Process', 'Assignment/Escalation/Auto-Response Rule', 'Duplicate/Matching Rule']),
    ('Quick Actions', 'F', ['Quick Action']),
    ('Reports / Dashboards / List Views', 'F', ['Report', 'Dashboard', 'List View']),
    ('Other Metadata', 'F', ['Other Metadata']),
    ('Page Layouts', 'U', ['Page Layout']),
    ('Lightning Pages', 'U', ['Lightning Page']),
    ('Compact Layouts / Path', 'U', ['Compact Layout', 'Path Assistant']),
    ('Record Types', 'C', ['Record Type']),
    ('Report Types', 'C', ['Report Type']),
    ('Permission Sets', 'C', ['Permission Set']),
    ('Profiles', 'C', ['Profile']),
    ('Translations / Field Sets', 'C', ['Translation', 'Field Set']),
]

def ncomp(api, cats, conf=None):
    n = 0
    for c in cats:
        for comp, v in R.get(api, {}).get(c, {}).items():
            if conf is None or v['confidence'].startswith(conf):
                n += 1
    return n

def fcount(api): return ncomp(api, FUNC)
def ucount(api): return ncomp(api, UI)
def ccount(api): return ncomp(api, CONF)
def probable(api): return ncomp(api, FUNC, 'Probable')

def comps(api, cats, limit=6):
    out = []
    for c in cats:
        for comp in sorted(R.get(api, {}).get(c, {})):
            out.append(comp)
    s = ', '.join(out[:limit])
    return s + (f' (+{len(out) - limit} more)' if len(out) > limit else '')

SHORT = {'Apex Class': 'Apex', 'Apex Trigger': 'Trigger', 'Flow': 'Flow', 'LWC': 'LWC', 'Aura': 'Aura',
         'Visualforce': 'VF', 'Email Template': 'Email', 'Validation Rule': 'Validation Rule',
         'Case Formula/Field': 'Case Field', 'Cross-Object Formula/Field': 'Other-Object Field',
         'Workflow': 'Workflow', 'Approval Process': 'Approval', 'Assignment/Escalation/Auto-Response Rule': 'Rule',
         'Quick Action': 'Quick Action', 'Report': 'Report', 'Dashboard': 'Dashboard', 'List View': 'List View',
         'Duplicate/Matching Rule': 'Dup Rule', 'Other Metadata': 'Other'}

def func_ref_text(api):
    parts = []
    for c in FUNC:
        cs = sorted(R.get(api, {}).get(c, {}))
        if cs:
            parts.append(f"{SHORT[c]} ({len(cs)}): " + ', '.join(cs[:5]) + (' …' if len(cs) > 5 else ''))
    return ' | '.join(parts) if parts else 'None found'

def layouts_text(api, limit=4):
    ls = sorted(R.get(api, {}).get('Page Layout', {}))
    ls = [re.sub(r'^Case-', '', x) for x in ls]
    s = '; '.join(ls[:limit])
    return (s + (f' (+{len(ls) - limit} more)' if len(ls) > limit else '')) if ls else ''

PREFIX = [('ccjo_', 'Credit Card – Jordan'), ('ccol_', 'Cash Collateral'), ('ppc_', 'Prepaid Card'), ('cc_', 'Credit Card'),
          ('AIR_', 'Account Interest Rate'), ('BUA_', 'Block / Unblock Account'), ('ac_', 'Account Closure'),
          ('bs_', 'Bond / Sukuk'), ('JO_', 'Jordan'), ('eKYC_', 'eKYC / Onboarding'), ('db_', 'Account Hold'),
          ('RHA_', 'Account Hold (RHA)'), ('sc_', 'Supplementary Card'), ('pc', 'Prepaid Card'), ('pr_', 'Profile / Onboarding')]
def module(api):
    if not F[api]['is_custom']:
        return 'Standard Salesforce'
    for p, m in PREFIX:
        if api.startswith(p):
            return m + ' (from API prefix)'
    lay = sorted(R.get(api, {}).get('Page Layout', {}))
    if lay:
        return re.sub(r'^Case-', '', lay[0]).replace(' Layout', '') + ' (from layout)'
    return 'General / Unknown'

# ---------------------------------------------------------------- pair & group classification
def toks(api): return re.sub(r'([a-z])([A-Z])', r'\1 \2', F[api]['label']).lower()
def tokset(api): return set(re.sub(r'[^a-z0-9 ]+', ' ', toks(api)).split()) - {'the', 'of', 'a', 'an', 'and'}

def mislabeled(api):
    f = F[api]
    return 'IMAGE(' in f['formula'].upper() and not re.search(r'logo|image|icon|flag|indicator', f['label'], re.I)

def classify(a, b, reason):
    if mislabeled(a) or mislabeled(b):
        return 'Related'
    if reason.startswith(('Identical', 'Same label words')):
        return 'Duplicate'
    if 'number/version suffix' in reason:
        return 'Numbered Series'
    A, B = tokset(a), tokset(b)
    x, y = A - B, B - A
    if len(x) == 1 and len(y) == 1:
        p, q = next(iter(x)), next(iter(y))
        if p[:4] == q[:4] and len(p) >= 4:
            return 'Duplicate'  # spelling variant e.g. Create / Created
        if re.fullmatch(r'l?\d+', p) and re.fullmatch(r'l?\d+', q):
            return 'Numbered Series'
    return 'Related'

pairs = defaultdict(list)
edge_rows = []
for a, b, reason, score in d['edges']:
    cls = classify(a, b, reason)
    edge_rows.append((a, b, reason, score, cls))
    pairs[a].append((b, reason, score, cls))
    pairs[b].append((a, reason, score, cls))

def keeper_score(api):
    f = F[api]
    return (0 if f['is_custom'] else 1, fcount(api) - probable(api) * 0.5, ucount(api),
            int(f['required']) + int(f['external_id']), ccount(api), -len(api))

CONTEXT_WORDS = ['residence', 'staff', 'update', 'jo', 'by', 'staff']
groups = []
for g in d['groups']:
    gcls = Counter(c for a in g for (b, _r, _s, c) in pairs[a] if b in g)
    cls = 'Duplicate' if gcls['Duplicate'] else ('Numbered Series' if gcls['Numbered Series'] and not gcls['Related'] else 'Related')
    groups.append(dict(members=g, cls=cls))

ORDER = {'Duplicate': 0, 'Related': 1, 'Numbered Series': 2}
groups.sort(key=lambda x: (ORDER[x['cls']], -len(x['members']), F[sorted(x['members'])[0]]['label'].lower()))

def group_name(g):
    labels = [F[a]['label'] for a in g]
    c = Counter(l.strip().lower() for l in labels).most_common(1)[0]
    if c[1] > 1:
        return next(l for l in labels if l.strip().lower() == c[0]).strip()
    # common leading words
    w = [l.split() for l in labels]
    common = []
    for col in zip(*w):
        if len({x.lower() for x in col}) == 1: common.append(col[0])
        else: break
    return (' '.join(common) + ' …') if common else labels[0]

field_rows = []
for gi, grp in enumerate(groups, 1):
    gid = f'G{gi:03d}'
    grp['id'] = gid
    g = grp['members']
    grp['name'] = group_name(g)
    keeper = None
    if grp['cls'] == 'Duplicate':
        dup_members = {a for a in g for (b, _, _, c) in pairs[a] if c == 'Duplicate' and b in g}
        keeper = max(dup_members, key=keeper_score)
    grp['keeper'] = keeper
    for api in sorted(g, key=lambda a: (a != keeper, -fcount(a), a)):
        f = F[api]
        fc, uc, cc, pr = fcount(api), ucount(api), ccount(api), probable(api)
        best = max([p for p in pairs[api] if p[0] in g], key=lambda p: p[2])
        dup_partners = [p for p in pairs[api] if p[0] in g and p[3] == 'Duplicate']
        type_note = ''
        if keeper and api != keeper and f['type'] != F[keeper]['type']:
            type_note = f" NOTE: data type differs ({f['type']} vs {F[keeper]['type']}) – map/convert values when repointing."
        ctx_note = ''
        if keeper and api != keeper:
            extra = set(re.findall(r'[a-z0-9]+', api.lower().replace('__c', '').replace('_', ' '))) - set(re.findall(r'[a-z0-9]+', keeper.lower().replace('__c', '').replace('_', ' ')))
            hits = [w for w in ('residence', 'staff', 'update', 'jo', 'new', 'old') if w in extra]
            if hits:
                ctx_note = f" API name contains '{', '.join(hits)}', which suggests a different business context – confirm with the process owner; if it is genuinely distinct, keep it and give it a unique label instead."
        is_f = f['is_formula']
        data_step = ('Formula field – no stored data to migrate.' if is_f else
                     f"Check data: SELECT COUNT() FROM Case WHERE {api} != null; if populated, back up and copy the values into {keeper} with Data Loader.")
        k = keeper
        if not f['is_custom']:
            role, prio = 'KEEP – Standard field', 'N/A'
            action = 'Standard Salesforce field – cannot be deleted. Use as the master field and retire the custom look-alike(s) in this group.'
        elif api == keeper:
            role, prio = 'KEEP – Master field', 'N/A'
            action = 'Retain as the single source of truth for this data. Point all references from the other field(s) in this group to this field.'
            if is_f:
                action += f" This is a formula ({f['formula'][:90]}…) – it is read-only, so it can only replace look-alikes that also just display this value."
        elif mislabeled(api):
            role, prio = 'KEEP – Rename label (mislabeled)', 'Low'
            action = (f"Its label says '{f['label']}', but the formula displays an image (IMAGE(…) – brand logo). It is not a duplicate of the similarly named field. "
                      f"Rename the label/description (e.g. 'Brand Logo') so nobody confuses it with a date/time field.")
        elif grp['cls'] == 'Duplicate' and dup_partners:
            if fc == 0 and uc == 0:
                role, prio = 'DELETE – Unused duplicate', 'High'
                action = (f"No code, automation, report or layout reference found (only profiles/permission sets/record types). 1) {data_step} "
                          f"2) Delete the field in Setup – permission-set, profile and report-type entries are removed automatically. 3) Remove it from the repo.")
            elif fc == 0:
                role, prio = 'DELETE – After layout cleanup', 'High'
                action = (f"Used only on UI pages: 1) Remove it from {uc} page layout(s)/Lightning page(s) ({comps(api, UI, 4)}) and add {k} where users need it. "
                          f"2) {data_step} 3) Delete the field.")
            elif F[k]['is_formula'] and not is_f:
                role, prio = 'REVIEW – Editable field vs formula master', 'Medium'
                action = (f"{k} is a read-only formula ({F[k]['formula'][:70]}…), while {api} is an editable field with {fc} functional reference(s) that may write to it "
                          f"({func_ref_text(api)}). It can be retired only if those processes can read the value from the formula's source instead of capturing it on the Case. "
                          f"If yes: repoint the references to {k}, remove it from layouts, back up data, delete. If no: keep it and give it a distinct label (e.g. add the product/region).")
            elif not F[k]['is_custom'] and f['type'] in ('Picklist', 'MultiselectPicklist'):
                role, prio = 'REVIEW – Compare with standard field', 'Low'
                action = (f"Look-alike of the standard field {k}. Compare picklist values: if the standard field's values can cover this process, move the values into {k}, "
                          f"repoint the {fc} reference(s) ({func_ref_text(api)}), migrate data and delete. If the values are process-specific, keep it and rename the label (e.g. '{f['label']} ({module(api).split(' (')[0]})').")
            elif fc >= 10 and fcount(k) >= 10:
                role, prio = 'REVIEW – Both heavily used', 'Low'
                action = (f"Both this field ({fc} functional refs) and {k} ({fcount(k)}) are heavily used. Confirm with the business that they hold the same data before consolidating. "
                          f"If confirmed, treat it as a mini-project: repoint every reference listed on the 'Functional References' sheet, migrate data, regression-test, then delete.")
            else:
                role, prio = 'CONSOLIDATE – Repoint then delete', 'Medium' if fc <= 5 else 'Low'
                action = (f"Has {fc} functional reference(s): {func_ref_text(api)}. 1) Replace {api} with {k} in each component (see 'Functional References' sheet) and deploy. "
                          f"2) Remove it from layouts/Lightning pages ({uc}). 3) {data_step} 4) Run Apex tests and regression-test the flows. 5) Delete the field.")
            action += type_note + ctx_note
        else:
            if fc == 0 and uc == 0:
                role, prio = 'DELETE – Unused (similar name)', 'High'
                action = (f"Not a true duplicate, but not referenced by any code, automation, report or page layout (only profiles/permission sets/record types). "
                          + ('Formula field – no data to back up; delete it.' if is_f else f"Check data with SELECT COUNT() FROM Case WHERE {api} != null; export the data if any, then delete the field."))
            elif fc == 0:
                role, prio = 'REVIEW – UI only', 'Medium'
                action = (f"Shown on {uc} layout(s)/page(s) ({comps(api, UI, 3)}) but not used by code, automation or reports. Confirm with users whether they still enter data here; "
                          f"if not, remove it from the layouts and delete it.")
            else:
                role, prio = 'KEEP – Distinct purpose', 'N/A'
                action = 'Name is similar to another field, but it stores different data (different stage, level, party or attribute) and is in use. Keep. Optionally improve the label/description so it is easier to tell apart.'
        derived = ''
        if is_f:
            srcs = [m for m in g if m != api and re.search(r'\b' + re.escape(m) + r'\b', f['formula'])]
            derived = (f"Derived from {', '.join(srcs)} – " if srcs else '') + 'Formula: ' + f['formula'][:200].replace('\n', ' ')
        bucket_vals = [ncomp(api, cats) for _, _, cats in BUCKETS]
        field_rows.append(dict(
            gid=gid, gname=grp['name'], gcls=grp['cls'], label=f['label'], api=api, type=f['type'] + (' (Formula)' if f['is_formula'] else ''),
            length=f['length'], custom='Custom' if f['is_custom'] else 'Standard', similar_to=f"{F[best[0]]['label']} ({best[0]})",
            basis=best[1], score=round(best[2], 2), module=module(api), layouts=layouts_text(api), description=(f['description'] or f['help'])[:300],
            flags=', '.join(x for x, y in [('Required', f['required']), ('External ID', f['external_id']), ('Unique', f['unique']), ('History Tracked', f['track_history'])] if y),
            buckets=bucket_vals, probable=pr, derived=derived, func_text=func_ref_text(api), role=role, prio=prio, action=action, keeper=keeper or '—'))

# ---------------------------------------------------------------- styling helpers
ARIAL = 'Arial'
HFILL = PatternFill('solid', start_color='1F3864')
HFONT = Font(name=ARIAL, bold=True, color='FFFFFF', size=10)
BFONT = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
TITLE = Font(name=ARIAL, size=16, bold=True, color='1F3864')
SUB = Font(name=ARIAL, size=10, italic=True, color='595959')
thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='top', wrap_text=True)
BAND = [PatternFill('solid', start_color='FFFFFF'), PatternFill('solid', start_color='EEF3FA')]
CATFILL = {'F': PatternFill('solid', start_color='2E75B6'), 'U': PatternFill('solid', start_color='548235'), 'C': PatternFill('solid', start_color='7F6000'),
           'T': PatternFill('solid', start_color='C00000')}

def header(ws, row, cols, widths=None, fills=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill, cell.alignment, cell.border = HFONT, (fills or {}).get(i, HFILL), Alignment(wrap_text=True, vertical='center', horizontal='center'), BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 45

def role_formats(ws, rng, col):
    rules = [('DELETE', 'F8CBAD', '833C0B'), ('CONSOLIDATE', 'FFE699', '7F6000'), ('REVIEW', 'FFF2CC', '7F6000'), ('KEEP', 'C6EFCE', '006100')]
    for key, bg, fg in rules:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${col}{rng.split(":")[0][len(col):]},{len(key)})="{key}"'],
                                                       fill=PatternFill('solid', start_color=bg, end_color=bg), font=Font(name=ARIAL, color=fg, bold=True)))

wb = Workbook()

# ================================================================ Field Details
ws = wb.active
ws.title = 'Field Details'
cols = (['Group ID', 'Group Name', 'Group Type', 'Field Label', 'API Name', 'Data Type', 'Length / Precision', 'Standard / Custom',
         'Most Similar Field', 'Similarity Basis', 'Similarity Score', 'Functional Area / Module', 'Page Layouts Using It', 'Description / Help Text', 'Field Flags', 'Formula / Derived From']
        + [b[0] for b in BUCKETS]
        + ['Total Functional Refs', 'Total UI Refs', 'Total Config Refs', 'Total References (Count)', 'Probable Refs (name shared w/ other objects)',
           'Usage Status', 'Functional Reference (components)', 'Recommended Master Field', 'Recommendation', 'Priority', 'Suggested Action', 'Owner / Decision', 'Status'])
nb = len(BUCKETS)
widths = [9, 26, 13, 30, 38, 14, 9, 10, 40, 30, 9, 24, 34, 40, 16, 40] + [11] * nb + [11, 10, 10, 12, 13, 15, 60, 34, 26, 9, 80, 16, 14]
fills = {}
for i, (_, k, _) in enumerate(BUCKETS):
    fills[17 + i] = CATFILL[k]
for j in range(4):
    fills[17 + nb + j] = CATFILL['T']
header(ws, 1, cols, widths, fills)
first_b = 17
colL = lambda i: get_column_letter(i)
fL = colL(first_b); fR = colL(first_b + 9)
uL = colL(first_b + 10); uR = colL(first_b + 12)
cL = colL(first_b + 13); cR = colL(first_b + nb - 1)
TF, TU, TC, TT = (colL(first_b + nb + i) for i in range(4))
PROB = colL(first_b + nb + 4); USAGE = colL(first_b + nb + 5)
ROLE = colL(first_b + nb + 8)
band, prev = 0, None
for r, fr in enumerate(field_rows, 2):
    if fr['gid'] != prev:
        band ^= 1; prev = fr['gid']
    vals = [fr['gid'], fr['gname'], fr['gcls'], fr['label'], fr['api'], fr['type'], fr['length'], fr['custom'], fr['similar_to'], fr['basis'], fr['score'],
            fr['module'], fr['layouts'], fr['description'], fr['flags'], fr['derived']] + fr['buckets'] + [
            f'=SUM({fL}{r}:{fR}{r})', f'=SUM({uL}{r}:{uR}{r})', f'=SUM({cL}{r}:{cR}{r})', f'={TF}{r}+{TU}{r}+{TC}{r}', fr['probable'],
            f'=IF(H{r}="Standard","Standard field",IF({TF}{r}>0,"In use",IF({TU}{r}>0,"UI only",IF({TC}{r}>0,"Security/config only","Unused"))))',
            fr['func_text'], fr['keeper'], fr['role'], fr['prio'], fr['action'], '', 'Not started']
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font, cell.border, cell.fill = BFONT, BORDER, BAND[band]
        cell.alignment = CENTER if (first_b <= c < first_b + nb + 5 or c in (1, 3, 7, 8, 11)) else WRAP
    ws.cell(row=r, column=11).number_format = '0%'
    ws.cell(row=r, column=5).font = Font(name='Consolas', size=10)
last = len(field_rows) + 1
ws.freeze_panes = 'F2'
ws.auto_filter.ref = f'A1:{colL(len(cols))}{last}'
role_formats(ws, f'{ROLE}2:{ROLE}{last}', ROLE)
dv = DataValidation(type='list', formula1='"Not started,In analysis,Approved,In progress,Deleted,Kept"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f'{colL(len(cols))}2:{colL(len(cols))}{last}')
FD = "'Field Details'"

# ================================================================ Duplicate Groups
wg = wb.create_sheet('Duplicate Groups', 0)
gcols = ['Group ID', 'Group Name', 'Group Type', 'Fields in Group (Count)', 'Field Labels', 'API Names', 'Data Types', 'Similarity Basis',
         'Recommended Master Field', 'Fields to Delete', 'Fields to Consolidate', 'Fields to Review', 'Fields to Keep', 'Total Functional Refs (group)',
         'Total References (group)', 'Group Recommendation']
header(wg, 1, gcols, [9, 30, 15, 11, 45, 50, 22, 40, 36, 10, 12, 10, 10, 12, 12, 70])
for r, grp in enumerate(groups, 2):
    g = grp['members']
    rows = [fr for fr in field_rows if fr['gid'] == grp['id']]
    bases = sorted({x[2] for x in edge_rows if x[0] in g and x[1] in g for x in [(x[0], x[1], x[2])]})
    roles = Counter(fr['role'].split(' ')[0] for fr in rows)
    if grp['cls'] == 'Duplicate':
        k = grp['keeper']
        rec = f"Proposed master: {k}. " + ('; '.join(f"{fr['api']}: {fr['role']}" for fr in rows if fr['api'] != k))
    elif any(fr['role'].startswith(('DELETE', 'REVIEW', 'KEEP – Rename')) for fr in rows):
        rec = 'Not true duplicates – similar names, different meaning. Clean up only the unused/UI-only members: ' + '; '.join(f"{fr['api']}: {fr['role']}" for fr in rows if not fr['role'].startswith(('KEEP – Distinct', 'KEEP – Master', 'KEEP – Standard')))
    else:
        rec = 'Not true duplicates – each field holds different data and is in use. No deletion needed; consider clearer labels/descriptions.'
    vals = [grp['id'], grp['name'], grp['cls'], f'=COUNTIF({FD}!$A:$A,A{r})', '\n'.join(F[a]['label'] for a in g), '\n'.join(g),
            '\n'.join(sorted({F[a]['type'] for a in g})), '\n'.join(bases), grp['keeper'] or '— (no master; not duplicates)',
            f'=COUNTIFS({FD}!$A:$A,A{r},{FD}!${ROLE}:${ROLE},"DELETE*")', f'=COUNTIFS({FD}!$A:$A,A{r},{FD}!${ROLE}:${ROLE},"CONSOLIDATE*")',
            f'=COUNTIFS({FD}!$A:$A,A{r},{FD}!${ROLE}:${ROLE},"REVIEW*")', f'=COUNTIFS({FD}!$A:$A,A{r},{FD}!${ROLE}:${ROLE},"KEEP*")',
            f'=SUMIF({FD}!$A:$A,A{r},{FD}!${TF}:${TF})', f'=SUMIF({FD}!$A:$A,A{r},{FD}!${TT}:${TT})', rec]
    for c, v in enumerate(vals, 1):
        cell = wg.cell(row=r, column=c, value=v)
        cell.font, cell.border, cell.alignment = BFONT, BORDER, (CENTER if c in (1, 3, 4, 10, 11, 12, 13, 14, 15) else WRAP)
    wg.cell(row=r, column=6).font = Font(name='Consolas', size=9)
glast = len(groups) + 1
wg.freeze_panes = 'C2'
wg.auto_filter.ref = f'A1:P{glast}'
for key, bg in [('Duplicate', 'F8CBAD'), ('Related', 'DDEBF7'), ('Numbered Series', 'E2EFDA')]:
    wg.conditional_formatting.add(f'C2:C{glast}', FormulaRule(formula=[f'$C2="{key}"'], fill=PatternFill('solid', start_color=bg, end_color=bg)))

# ================================================================ Functional References
wr = wb.create_sheet('Functional References')
rcols = ['Group ID', 'Field Label', 'API Name', 'Recommendation', 'Reference Category', 'Reference Type', 'Component Name', 'Occurrences', 'Confidence', 'File Path(s)', 'Change Needed']
header(wr, 1, rcols, [9, 30, 38, 26, 14, 26, 42, 11, 30, 70, 60])
r = 2
kind = {**{c: 'Functional' for c in FUNC}, **{c: 'UI' for c in UI}, **{c: 'Security / Config' for c in CONF}}
for fr in field_rows:
    api = fr['api']
    for cat in FUNC + UI + CONF:
        if fr['role'].startswith('KEEP') and kind[cat] == 'Security / Config':
            continue
        for comp, v in sorted(R.get(api, {}).get(cat, {}).items()):
            if fr['role'].startswith('KEEP'):
                change = 'None – field is kept'
            elif kind[cat] == 'Security / Config':
                change = 'Removed automatically by Salesforce when the field is deleted; also remove from the repo source.'
            elif kind[cat] == 'UI':
                change = f"Remove field from this {cat.lower()}" + (f" and add {fr['keeper']} if needed" if fr['keeper'] != '—' else '')
            else:
                change = (f"Replace {api} with {fr['keeper']}" if fr['keeper'] != '—' else f'Remove reference to {api}') + ' and redeploy'
            vals = [fr['gid'], fr['label'], api, fr['role'], kind[cat], cat, comp, v['hits'], v['confidence'], '\n'.join(v['files'][:5]), change]
            for c, val in enumerate(vals, 1):
                cell = wr.cell(row=r, column=c, value=val)
                cell.font, cell.border, cell.alignment = BFONT, BORDER, (CENTER if c in (1, 5, 8) else WRAP)
            r += 1
rlast = r - 1
wr.freeze_panes = 'D2'
wr.auto_filter.ref = f'A1:K{rlast}'
role_formats(wr, f'D2:D{rlast}', 'D')

# ================================================================ All Case Fields inventory
wa = wb.create_sheet('All Case Fields Usage')
acols = ['Field Label', 'API Name', 'Data Type', 'Standard / Custom', 'In Similar-Name Group', 'Functional Area / Module', 'Functional Refs', 'UI Refs', 'Config Refs',
         'Total References (Count)', 'Usage Status', 'Functional Reference (components)', 'Description']
header(wa, 1, acols, [32, 40, 16, 11, 12, 26, 11, 9, 9, 12, 18, 70, 50])
gid_of = {fr['api']: fr['gid'] for fr in field_rows}
for r, api in enumerate(sorted(F, key=lambda a: F[a]['label'].lower()), 2):
    f = F[api]
    std = not f['is_custom']
    vals = [f['label'], api, f['type'] + (' (Formula)' if f['is_formula'] else ''), 'Standard' if std else 'Custom', gid_of.get(api, ''), module(api),
            None if std else fcount(api), None if std else ucount(api), None if std else ccount(api),
            None if std else f'=G{r}+H{r}+I{r}',
            'Standard field' if std else f'=IF(G{r}>0,"In use",IF(H{r}>0,"UI only",IF(I{r}>0,"Security/config only","Unused")))',
            'Not analysed (standard field)' if std else func_ref_text(api), f['description'][:250]]
    for c, v in enumerate(vals, 1):
        cell = wa.cell(row=r, column=c, value=v)
        cell.font, cell.border, cell.alignment = BFONT, BORDER, (CENTER if c in (4, 5, 7, 8, 9, 10, 11) else WRAP)
alast = len(F) + 1
wa.freeze_panes = 'C2'
wa.auto_filter.ref = f'A1:M{alast}'
for key, bg in [('Unused', 'F8CBAD'), ('Security/config only', 'FCE4D6'), ('UI only', 'FFF2CC'), ('In use', 'C6EFCE')]:
    wa.conditional_formatting.add(f'K2:K{alast}', FormulaRule(formula=[f'$K2="{key}"'], fill=PatternFill('solid', start_color=bg, end_color=bg)))

# ================================================================ Summary (first sheet)
wsu = wb.create_sheet('Summary', 0)
wsu.column_dimensions['A'].width = 4
wsu.column_dimensions['B'].width = 48
wsu.column_dimensions['C'].width = 14
wsu.column_dimensions['D'].width = 90
wsu['B2'] = 'Case Object – Similar / Duplicate Field Analysis'; wsu['B2'].font = TITLE
wsu['B3'] = 'Source: Salesforce metadata in force-app/main/default (repository UAT-Project-06Oct2026), analysed 11-Oct-2026. All counts below are live formulas over the detail sheets.'
wsu['B3'].font = SUB
AF = "'All Case Fields Usage'"; DG = "'Duplicate Groups'"
rows = [
    ('Field inventory', None, None),
    ('Total Case fields in metadata', f'=COUNTA({AF}!B2:B{alast})', 'Standard + custom fields found under objects/Case/fields'),
    ('Custom Case fields', f'=COUNTIF({AF}!D2:D{alast},"Custom")', 'Only custom fields can be deleted'),
    ('Custom fields with no reference at all (Unused)', f'=COUNTIF({AF}!K2:K{alast},"Unused")', "Not on any layout, code, flow, report, permission set… – see 'All Case Fields Usage'"),
    ('Custom fields referenced only by profiles / permission sets / record types', f'=COUNTIF({AF}!K2:K{alast},"Security/config only")', 'No business usage – strong deletion candidates after a data check'),
    ('Custom fields used only on page layouts', f'=COUNTIF({AF}!K2:K{alast},"UI only")', 'Shown to users but not used by any automation, code or report'),
    ('Similar-name groups', None, None),
    ('Groups of fields with similar label / API name', f'=COUNTA({DG}!A2:A{glast})', 'Fields matched on label and API name (exact, word order, version suffix, fuzzy)'),
    ('   – Duplicate groups (same data stored twice)', f'=COUNTIF({DG}!C2:C{glast},"Duplicate")', 'Consolidate into one master field'),
    ('   – Related groups (similar name, different meaning)', f'=COUNTIF({DG}!C2:C{glast},"Related")', 'e.g. Result vs Result Date; usually keep'),
    ('   – Numbered series (L2/L3, 1/2/3…)', f'=COUNTIF({DG}!C2:C{glast},"Numbered Series")', 'Intentional series; keep unless unused'),
    ('Fields in similar-name groups', f'=COUNTA({FD}!E2:E{last})', None),
    ('Recommendations for grouped fields', None, None),
    ('DELETE – no blocking references (quick wins)', f'=COUNTIF({FD}!{ROLE}2:{ROLE}{last},"DELETE*")', 'Data check, then delete. Priority High'),
    ('CONSOLIDATE – repoint references, migrate data, then delete', f'=COUNTIF({FD}!{ROLE}2:{ROLE}{last},"CONSOLIDATE*")', 'Needs development / config changes. Priority Medium/Low'),
    ('REVIEW – business decision required', f'=COUNTIF({FD}!{ROLE}2:{ROLE}{last},"REVIEW*")', 'UI-only or both fields heavily used'),
    ('KEEP', f'=COUNTIF({FD}!{ROLE}2:{ROLE}{last},"KEEP*")', 'Master fields and fields with a distinct purpose'),
    ('Functional references to be repointed before deletion', f'=SUMIFS({FD}!{TF}2:{TF}{last},{FD}!{ROLE}2:{ROLE}{last},"CONSOLIDATE*")+SUMIFS({FD}!{TF}2:{TF}{last},{FD}!{ROLE}2:{ROLE}{last},"REVIEW*")', 'Apex, flows, reports, validation rules… on CONSOLIDATE / REVIEW fields'),
]
r = 5
for label, formula, note in rows:
    if formula is None and note is None:
        r += 1
        c = wsu.cell(row=r, column=2, value=label); c.font = Font(name=ARIAL, bold=True, size=11, color='FFFFFF'); c.fill = HFILL
        wsu.cell(row=r, column=3).fill = HFILL; wsu.cell(row=r, column=4).fill = HFILL
    else:
        wsu.cell(row=r, column=2, value=label).font = BFONT
        c = wsu.cell(row=r, column=3, value=formula); c.font = BOLD; c.alignment = Alignment(horizontal='center')
        wsu.cell(row=r, column=4, value=note).font = SUB
        for cc in (2, 3, 4): wsu.cell(row=r, column=cc).border = BORDER
    r += 1
r += 1
wsu.cell(row=r, column=2, value='How to read this workbook').font = Font(name=ARIAL, bold=True, size=11, color='1F3864'); r += 1
guide = [
    ('Duplicate Groups', 'One row per group of similar fields, with the recommended master field and counts by recommendation.'),
    ('Field Details', 'One row per field in a group: label, API name, type, similarity, reference counts by category, usage status, recommendation, priority and step-by-step action. Columns "Owner / Decision" and "Status" are for you to fill in.'),
    ('Functional References', 'Every component (Apex class, flow, report, layout, permission set…) that references a grouped field, and the change needed before deletion.'),
    ('All Case Fields Usage', 'All Case fields with reference counts – use it to find unused fields beyond the duplicate groups.'),
    ('Deletion Playbook', 'The ordered steps to safely delete a field in Salesforce.'),
    ('Methodology', 'How similarity and usage were determined, and the limits of this analysis.'),
]
for s, t in guide:
    wsu.cell(row=r, column=2, value=s).font = BOLD
    c = wsu.cell(row=r, column=4, value=t); c.font = BFONT; c.alignment = WRAP
    r += 1
r += 1
wsu.cell(row=r, column=2, value='Colour legend (Recommendation column)').font = Font(name=ARIAL, bold=True, size=11, color='1F3864'); r += 1
for name, bg, txt_ in [('DELETE', 'F8CBAD', 'Safe to delete after a data check'), ('CONSOLIDATE', 'FFE699', 'Repoint references to the master field, migrate data, then delete'),
                       ('REVIEW', 'FFF2CC', 'Needs a business decision'), ('KEEP', 'C6EFCE', 'Keep the field')]:
    c = wsu.cell(row=r, column=2, value=name); c.font = BOLD; c.fill = PatternFill('solid', start_color=bg)
    wsu.cell(row=r, column=4, value=txt_).font = BFONT
    r += 1

# ================================================================ Deletion Playbook
wp = wb.create_sheet('Deletion Playbook')
header(wp, 1, ['Step', 'Phase', 'Action', 'How / Tool', 'Applies To'], [6, 18, 60, 70, 26])
steps = [
    ('Prepare', 'Confirm the master field for each Duplicate group with the business / process owner.', "Review 'Duplicate Groups'. Record the decision in 'Field Details' → Owner / Decision.", 'All CONSOLIDATE / REVIEW'),
    ('Prepare', 'Check whether each field to delete still holds data.', 'Developer Console / Workbench: SELECT COUNT() FROM Case WHERE <API_Name> != null (use != false for checkboxes). Also check last-populated date: SELECT MAX(LastModifiedDate) FROM Case WHERE <API_Name> != null.', 'All DELETE / CONSOLIDATE'),
    ('Prepare', 'Check the org for references that are not in source control.', "Setup → Object Manager → Case → Fields → <field> → 'Where is this used?'. Also check managed packages, external integrations (Middleware, ETL, APIs) and Einstein / CRM Analytics datasets.", 'All fields to delete'),
    ('Prepare', 'Back up the data.', 'Data Loader / Data Export: export Id + field values for every field that will be deleted.', 'All fields to delete'),
    ('Migrate', 'Copy data into the master field when the field to delete holds values the master does not.', 'Data Loader update, or a one-off Apex batch. Map picklist values / convert data types where the types differ (see NOTE in Suggested Action).', 'CONSOLIDATE (and DELETE with data)'),
    ('Repoint', 'Replace references to the field with the master field.', "Use 'Functional References' sheet: update Apex classes/triggers, flows, LWC/Aura/VF, email templates, validation rules, formula fields, workflow/approval processes, quick actions, reports, dashboards and list views. Deploy via the normal release pipeline.", 'CONSOLIDATE'),
    ('Repoint', 'Remove the field from page layouts, Lightning record pages, compact layouts and path.', 'Edit layouts in Setup or remove the <field> entries from the layout/flexipage XML in the repo.', 'DELETE – After layout cleanup / CONSOLIDATE'),
    ('Test', 'Regression-test in a sandbox (UAT).', 'Run all Apex tests (sf apex run test --test-level RunLocalTests), retest affected flows and record types, rerun affected reports.', 'All'),
    ('Delete', 'Delete the field in the sandbox, then production.', "Setup → Object Manager → Case → Fields & Relationships → Delete, or deploy a destructiveChanges.xml (CustomField: Case.<API_Name>) with the sf CLI. Profile, permission-set, report-type and record-type entries are removed automatically.", 'All fields to delete'),
    ('Delete', 'Remove the field from source control.', 'Delete objects/Case/fields/<API_Name>.field-meta.xml and references in permission sets, profiles, layouts, record types and translations; remove from manifest/package.xml.', 'All fields to delete'),
    ('After', 'Erase deleted fields permanently when ready.', 'Deleted custom fields stay recoverable for 15 days and still count against the field limit. Use Setup → Deleted Fields → Erase to free the slot immediately (cannot be undone).', 'All deleted fields'),
    ('After', 'Update documentation and the tracking status.', "Set 'Status' to Deleted / Kept in 'Field Details'. Update the data dictionary.", 'All'),
]
for r, (ph, a, h, ap) in enumerate(steps, 2):
    for c, v in enumerate([r - 1, ph, a, h, ap], 1):
        cell = wp.cell(row=r, column=c, value=v)
        cell.font, cell.border, cell.alignment = BFONT, BORDER, (CENTER if c == 1 else WRAP)
wp.freeze_panes = 'A2'

# ================================================================ Methodology
wm = wb.create_sheet('Methodology')
wm.column_dimensions['A'].width = 30; wm.column_dimensions['B'].width = 120
items = [
    ('Scope', f"All {len(F)} Case fields defined in force-app/main/default/objects/Case/fields ({sum(1 for f in F.values() if f['is_custom'])} custom, {sum(1 for f in F.values() if not f['is_custom'])} standard). Only custom fields are considered for deletion."),
    ('Label similarity', 'Labels are lower-cased, punctuation removed and camel-case split. Matches: identical label; same words in a different order; label differing only by a trailing number/version word (2, v2, new, old, copy…); one label equal to the other plus a single qualifier word (brand, region, stage or role – e.g. "Alburaq Customer Segment" vs "Customer Segment", "Approved Credit Limit" vs "Credit Limit"); or fuzzy similarity ≥ 90% (Python difflib). Qualifier matches are treated as Related (different meaning) unless other evidence shows the same data.'),
    ('API name similarity', 'API names are compared ignoring case, underscores and the __c suffix. Matches: identical; differing only by a trailing number/version word; or fuzzy similarity ≥ 92%.'),
    ('Grouping', 'Matched pairs are linked into groups (connected components). A group is "Duplicate" when at least one pair has the same meaning (identical label/API, or a spelling variant such as Create/Created); "Numbered Series" when members differ only by level/number (L2/L3, 1/2); otherwise "Related" (similar name, different data).'),
    ('Master field', 'Within a Duplicate group: a standard field always wins; otherwise the field with the most confirmed functional references, then most UI references, then required/external-ID flags.'),
    ('Usage / references', 'Every text file under force-app/main/default (≈28,000 files) was scanned for the field API name. References are counted per component (one Apex class = 1 reference, no matter how many times it mentions the field). Categories: Functional (Apex, triggers, flows, LWC, Aura, Visualforce, email templates, validation rules, formula fields, workflow, approval, assignment/escalation rules, quick actions, reports, dashboards, list views); UI (page layouts, Lightning pages, compact layouts, path); Security/Config (permission sets, profiles, record types, report types, translations).'),
    ('Confidence', 'A reference is "Confirmed" when it is in Case-specific metadata, is qualified with Case (e.g. Case.Field__c, Case__r.Field__c), or the API name exists only on Case. When the same API name also exists on another object and the file is generic (e.g. Apex), the reference is counted only if the file mentions Case, and is marked "Probable".'),
    ('Usage status', 'In use = at least one functional reference; UI only = only on layouts/pages; Security/config only = only in profiles, permission sets, record types or report types; Unused = no reference found.'),
    ('Limitations', 'This is a static analysis of the repository. It cannot see: field data population (run the SOQL in the Suggested Action), components not retrieved into the repo, managed packages, external integrations/ETL, CRM Analytics, or dynamic Apex that builds field names as strings. Always confirm with "Where is this used?" in Setup and a data check before deleting.'),
    ('Functional Area / Module', "Inferred from the API name prefix (cc_ = Credit Card, ccjo_ = Credit Card Jordan, ppc_ = Prepaid Card, ccol_ = Cash Collateral, AIR_ = Account Interest Rate, BUA_ = Block/Unblock Account, ac_ = Account Closure, bs_ = Bond/Sukuk, …) or, when there is no prefix, from the first page layout that uses the field. Prefix meanings are assumptions based on naming and record types – please validate."),
]
wm['A1'] = 'Methodology & Assumptions'; wm['A1'].font = TITLE
for r, (k, v) in enumerate(items, 3):
    a = wm.cell(row=r, column=1, value=k); a.font = BOLD; a.alignment = WRAP; a.border = BORDER
    b = wm.cell(row=r, column=2, value=v); b.font = BFONT; b.alignment = WRAP; b.border = BORDER

wb.move_sheet('Summary', -(wb.index(wb['Summary'])))
for s in wb.worksheets:
    s.sheet_view.zoomScale = 90
wb.save(OUT)
print('groups', len(groups), Counter(g['cls'] for g in groups), 'fields', len(field_rows), Counter(fr['role'] for fr in field_rows), 'ref rows', rlast - 1)
